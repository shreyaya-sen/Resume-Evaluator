import os
import time
import json
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
import gradio as gr

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client = Groq(api_key=my_api_key)
model = "qwen/qwen3.8-27b"


class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


jobd_schema = JobD.model_json_schema()


def parse_job_description(job_description):

    system_prompt = f"""
    You are an expert HR assistant.

    Your job is to analyze job descriptions and extract
    structured information from them.

    Return ONLY valid JSON matching this schema:

    {jobd_schema}

    IMPORTANT:
    Do NOT return the schema itself.
    Do NOT return fields like "properties", "title" or "type".
    Fill the schema with actual information extracted from the job description.

    If minimum experience is not mentioned, return null.
    If information for a list is missing, return an empty list.
    Do not invent information.
    """

    user_prompt = f"""
    Analyze the following job description:

    {job_description}
    """

    message_system = {
        "role": "system",
        "content": system_prompt
    }

    message_user = {
        "role": "user",
        "content": user_prompt
    }

    messages = [message_system, message_user]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format,
        max_completion_tokens=400
    )

    raw_json = response.choices[0].message.content

    job_data = json.loads(raw_json)

    job = JobD(**job_data)

    return job


class MatchResult(BaseModel):
    candidate_name: str
    score: float
    matching_skills: list[str]
    missing_skills: list[str]
    experience_requirement_met: bool | None
    assessment: str


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


def evaluate_single_resume(resume_text, job):

    match_schema = MatchResult.model_json_schema()

    prompt = f"""
    You are an expert HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume_text}

    First understand the resume, then compare it with the job description.

    Return ONLY valid JSON matching this schema:

    {match_schema}

    Requirements:

    1. Extract the candidate's name.
    2. Identify matching skills.
    3. Identify important missing skills.
    4. Determine whether the experience requirement is met.
    5. Give an overall match score from 0 to 100.
    6. Give a short professional assessment.
    7. Do not invent information.
    8. Keep the assessment concise.
    """

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        response_format={
            "type": "json_object"
        },
        max_completion_tokens=500
    )

    data = json.loads(
        response.choices[0].message.content
    )

    return MatchResult(**data)


from pypdf import PdfReader
from docx import Document


def read_pdf(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def read_docx(file_path):

    document = Document(file_path)

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_path):

    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)

    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)

    else:
        return None


def process_single_resume(file_path, job):

    file_path = Path(file_path)

    resume_text = read_resume(file_path)

    if not resume_text:
        return None

    result = evaluate_single_resume(resume_text, job)

    return {
        "name": result.candidate_name,
        "score": result.score,
        "assessment": result.assessment
    }


def evaluate_resumes(job_description, uploaded_files):
    if len(uploaded_files) > 3:
        raise gr.Error("Please upload a maximum of 3 resumes per evaluation.")
    
    job = parse_job_description(job_description)

    all_results = []

    with ThreadPoolExecutor(max_workers=2) as executor:

        futures = [
            executor.submit(
                process_single_resume,
                file_path,
                job
            )
            for file_path in uploaded_files
        ]

        for future in as_completed(futures):

            candidate = future.result()

            if candidate:
                all_results.append(candidate)

    all_results.sort(
        key=lambda candidate: candidate["score"],
        reverse=True
    )

    results = []

    for candidate in all_results:

        results.append([
            candidate["name"],
            f"{candidate['score']}%",
            candidate["assessment"]
        ])

    return results


with gr.Blocks(fill_width=True) as demo:

    gr.Markdown("# Resume Evaluator")

    with gr.Row():

        job_description_input = gr.Textbox(
            label="Job Description",
            placeholder="Paste the job description here...",
            lines=10,
            scale=1
        )

        resumes = gr.Files(
            label="Upload Resumes",
            file_types=[".pdf", ".docx"],
            type="filepath",
            scale=1
        )

    evaluate_button = gr.Button("Evaluate Resumes")

    results = gr.Dataframe(
        headers=["Candidate Name", "Match Score", "Assessment"],
        datatype=["str", "str", "str"],
        label="Evaluation Results",
        wrap=True,
        column_widths=["20%", "15%", "65%"],
        max_height=600
    )

    evaluate_button.click(
        fn=evaluate_resumes,
        inputs=[job_description_input, resumes],
        outputs=[results]
    )


demo.launch(
    server_name="0.0.0.0",
    server_port=int(os.environ.get("PORT", 10000)),
    show_error=True
)
